# -*- coding: utf-8 -*-

"""
============================================================
HWP Auto DocFit - 폴더 일괄 변환 (HWP/HWPX → PDF, XLSX → Markdown)
============================================================

사용자가 지정한 폴더(하위 폴더 포함) 안의

    - .hwp / .hwpx 문서는 같은 폴더에 동일 파일명의 .pdf로
    - .xlsx 통합 문서는 같은 폴더에 동일 파일명의 .md로

일괄 변환한다.

HWP/HWPX → PDF는 한/글 OLE Automation(pyhwpx)이 필요하므로
Windows + 한글 2020 이상 환경에서만 동작한다.
XLSX → Markdown은 openpyxl만으로 동작하며 운영체제를 가리지 않는다.

CLI:
    python folder_convert.py <폴더경로> [--no-recursive]

MCP 도구로도 노출된다: mcp_server.py의 convert_folder_documents.
============================================================
"""

import os
import sys
from pathlib import Path

try:
    from pyhwpx import Hwp
except ImportError:
    Hwp = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


HWP_EXTENSIONS = (".hwp", ".hwpx")
EXCEL_EXTENSIONS = (".xlsx",)


# ============================================================
# 대상 파일 탐색
# ============================================================

def find_target_files(folder_path: str, recursive: bool = True):
    """폴더 내 HWP/HWPX 파일과 XLSX 파일 목록을 각각 반환한다.

    Excel/한글이 열어둔 잠금 임시 파일(~$로 시작)은 제외한다.
    """
    folder = Path(folder_path)
    pattern = "**/*" if recursive else "*"

    hwp_files = []
    xlsx_files = []

    for p in sorted(folder.glob(pattern)):
        if not p.is_file() or p.name.startswith("~$"):
            continue
        suffix = p.suffix.lower()
        if suffix in HWP_EXTENSIONS:
            hwp_files.append(p)
        elif suffix in EXCEL_EXTENSIONS:
            xlsx_files.append(p)

    return hwp_files, xlsx_files


# ============================================================
# HWP / HWPX → PDF
# ============================================================

def convert_hwp_to_pdf(hwp_obj, src_path: Path) -> Path:
    """단일 HWP/HWPX 파일을 같은 폴더에 동일 파일명의 PDF로 저장한다."""
    dst_path = src_path.with_suffix(".pdf")
    hwp_obj.open(str(src_path))
    try:
        hwp_obj.save_as(str(dst_path), Format="PDF")
    finally:
        try:
            hwp_obj.clear(1)
        except Exception:
            pass

    if not dst_path.is_file():
        raise RuntimeError("PDF 파일 생성 확인 실패")

    return dst_path


# ============================================================
# XLSX → Markdown
# ============================================================

def _cell_to_md(value) -> str:
    if value is None:
        return ""
    text = str(value)
    # 표 구분자·줄바꿈은 Markdown 표를 깨뜨리므로 이스케이프/치환한다.
    text = text.replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>")
    return text


def sheet_to_markdown(ws) -> str:
    """openpyxl 워크시트 하나를 Markdown 표 문자열로 변환한다."""
    rows = list(ws.iter_rows(values_only=True))

    # 시트 끝의 완전히 빈 행은 잘라낸다 (openpyxl이 종종 넉넉하게 잡는 사용 범위 보정).
    while rows and all(cell is None for cell in rows[-1]):
        rows.pop()

    if not rows:
        return "_(빈 시트)_\n"

    col_count = max(len(r) for r in rows)

    def row_to_md(row):
        cells = [_cell_to_md(row[i]) if i < len(row) else "" for i in range(col_count)]
        return "| " + " | ".join(cells) + " |"

    lines = [row_to_md(rows[0]), "| " + " | ".join(["---"] * col_count) + " |"]
    lines.extend(row_to_md(r) for r in rows[1:])
    return "\n".join(lines) + "\n"


def convert_xlsx_to_md(src_path: Path) -> Path:
    """단일 XLSX 파일을 같은 폴더에 동일 파일명의 Markdown으로 저장한다.

    시트가 여러 개면 시트마다 '## 시트명' 제목으로 구분한다.
    """
    if openpyxl is None:
        raise RuntimeError("openpyxl이 설치되어 있지 않습니다. `pip install openpyxl`로 설치하세요.")

    dst_path = src_path.with_suffix(".md")
    wb = openpyxl.load_workbook(str(src_path), data_only=True, read_only=True)

    try:
        parts = [f"# {src_path.stem}\n"]
        multi_sheet = len(wb.sheetnames) > 1
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if multi_sheet:
                parts.append(f"## {sheet_name}\n")
            parts.append(sheet_to_markdown(ws))
    finally:
        wb.close()

    dst_path.write_text("\n".join(parts), encoding="utf-8")
    return dst_path


# ============================================================
# 폴더 일괄 변환
# ============================================================

def convert_folder(folder_path: str, recursive: bool = True, progress_cb=None) -> dict:
    """
    지정된 폴더의 모든 HWP/HWPX -> PDF, XLSX -> Markdown을 일괄 변환한다.

    progress_cb(message: str)가 주어지면 각 단계마다 호출해 진행 상황을 알린다.
    """
    if not os.path.isdir(folder_path):
        raise ValueError(f"폴더를 찾을 수 없습니다: {folder_path}")

    def log(message: str):
        if progress_cb:
            progress_cb(message)

    hwp_files, xlsx_files = find_target_files(folder_path, recursive=recursive)

    results = {
        "folder_path": folder_path,
        "pdf": [],
        "markdown": [],
        "errors": [],
    }

    # ------------------------------------------------------
    # HWP/HWPX -> PDF
    # ------------------------------------------------------

    if hwp_files:
        if Hwp is None:
            message = (
                "pyhwpx(win32com)를 사용할 수 없어 HWP/HWPX -> PDF 변환을 "
                "건너뜁니다. (Windows + 한글 2020 이상 환경에서만 동작)"
            )
            log(message)
            results["errors"].append({"type": "hwp_engine_missing", "message": message})
        else:
            hwp_obj = None
            try:
                log("한글 인스턴스 시작 중...")
                hwp_obj = Hwp(visible=False)
                try:
                    hwp_obj.RegisterModule("FilePathCheckDLL", "FilePathCheckerModule")
                except Exception:
                    pass

                for src in hwp_files:
                    try:
                        log(f"PDF 변환 중: {src.name}")
                        dst = convert_hwp_to_pdf(hwp_obj, src)
                        results["pdf"].append({"source": str(src), "output": str(dst)})
                        log(f"완료: {dst.name}")
                    except Exception as e:
                        log(f"실패: {src.name} ({e})")
                        results["errors"].append(
                            {"type": "hwp_to_pdf", "file": str(src), "message": str(e)}
                        )
            finally:
                if hwp_obj is not None:
                    try:
                        hwp_obj.quit()
                    except Exception:
                        pass

    # ------------------------------------------------------
    # XLSX -> Markdown
    # ------------------------------------------------------

    for src in xlsx_files:
        try:
            log(f"Markdown 변환 중: {src.name}")
            dst = convert_xlsx_to_md(src)
            results["markdown"].append({"source": str(src), "output": str(dst)})
            log(f"완료: {dst.name}")
        except Exception as e:
            log(f"실패: {src.name} ({e})")
            results["errors"].append(
                {"type": "xlsx_to_md", "file": str(src), "message": str(e)}
            )

    results["total_hwp_found"] = len(hwp_files)
    results["total_xlsx_found"] = len(xlsx_files)
    results["pdf_converted_count"] = len(results["pdf"])
    results["markdown_converted_count"] = len(results["markdown"])
    return results


# ============================================================
# CLI
# ============================================================

def cli_main(argv):
    import argparse

    parser = argparse.ArgumentParser(
        description="지정한 폴더의 HWP/HWPX -> PDF, XLSX -> Markdown 일괄 변환"
    )
    parser.add_argument("folder", help="변환할 폴더 경로")
    parser.add_argument(
        "--no-recursive",
        action="store_true",
        help="하위 폴더를 검색하지 않음 (기본값: 하위 폴더 포함)",
    )
    args = parser.parse_args(argv)

    print("=" * 70)
    print("HWP Auto DocFit - 폴더 일괄 변환")
    print(f"대상 폴더: {args.folder}")
    print("=" * 70)

    result = convert_folder(args.folder, recursive=not args.no_recursive, progress_cb=print)

    print()
    print("=" * 70)
    print(
        f"변환 완료: PDF {result['pdf_converted_count']}/{result['total_hwp_found']}개, "
        f"Markdown {result['markdown_converted_count']}/{result['total_xlsx_found']}개"
    )
    if result["errors"]:
        print(f"오류 {len(result['errors'])}건 발생 (상세는 위 로그 참고)")
    print("=" * 70)

    return result


if __name__ == "__main__":
    cli_main(sys.argv[1:])
