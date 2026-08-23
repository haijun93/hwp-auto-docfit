# -*- coding: utf-8 -*-

"""
============================================================
HWP Auto DocFit - folder_convert.py 단위 테스트
============================================================
XLSX -> Markdown 변환은 순수 파이썬(openpyxl)이라 어떤 OS에서도
실행 가능하다. HWP/HWPX -> PDF 변환은 pyhwpx(Windows + 한글)가
필요하므로, 엔진이 없는 환경(CI 등)에서는 "건너뛰고 오류로 보고"하는
동작만 검증한다.
============================================================
"""

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import openpyxl  # noqa: E402

from folder_convert import (  # noqa: E402
    convert_folder,
    convert_xlsx_to_md,
    find_target_files,
    sheet_to_markdown,
)


def run_all_tests():
    print("=" * 70)
    print("folder_convert.py 단위 테스트")
    print("=" * 70)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        sub = tmp_path / "sub"
        sub.mkdir()

        # ----------------------------------------------------
        # 1. 대상 파일 탐색 (재귀 + 잠금파일 제외)
        # ----------------------------------------------------

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "매출"
        ws1.append(["이름", "부서", "비고"])
        ws1.append(["홍길동", "기획팀", None])
        ws1.append(["김철수", "개발|팀", "여러 줄\n메모"])
        wb.create_sheet("현황").append(["A", "B"])
        wb.save(str(tmp_path / "data.xlsx"))

        wb2 = openpyxl.Workbook()
        wb2.active.append(["x", "y"])
        wb2.save(str(sub / "nested.xlsx"))

        (tmp_path / "~$data.xlsx").write_text("lock")
        (tmp_path / "doc.hwp").write_bytes(b"fake-hwp-bytes")

        hwp_files, xlsx_files = find_target_files(str(tmp_path), recursive=True)
        assert [f.name for f in hwp_files] == ["doc.hwp"], hwp_files
        assert sorted(f.name for f in xlsx_files) == ["data.xlsx", "nested.xlsx"], xlsx_files
        print("[TEST 1/4] 대상 파일 탐색 (재귀 포함, ~$ 잠금파일 제외): PASS")

        hwp_files_flat, xlsx_files_flat = find_target_files(str(tmp_path), recursive=False)
        assert [f.name for f in xlsx_files_flat] == ["data.xlsx"], xlsx_files_flat
        print("[TEST 2/4] --no-recursive 시 하위 폴더 미탐색: PASS")

        # ----------------------------------------------------
        # 2. 시트 -> Markdown 표 변환 (이스케이프 포함)
        # ----------------------------------------------------

        md = sheet_to_markdown(ws1)
        lines = md.strip().split("\n")
        assert lines[0] == "| 이름 | 부서 | 비고 |", lines[0]
        assert lines[1] == "| --- | --- | --- |", lines[1]
        assert "개발\\|팀" in lines[3], lines[3]   # '|' 이스케이프
        assert "<br>" in lines[3], lines[3]        # 줄바꿈 -> <br>
        print("[TEST 3/4] 시트 -> Markdown 표 변환 (파이프/줄바꿈 이스케이프): PASS")

        # ----------------------------------------------------
        # 3. 폴더 일괄 변환: XLSX는 성공, HWP는 엔진 없으면 오류로 보고
        # ----------------------------------------------------

        result = convert_folder(str(tmp_path), recursive=True)
        assert result["total_xlsx_found"] == 2, result
        assert result["markdown_converted_count"] == 2, result
        assert (tmp_path / "data.md").is_file()
        assert (sub / "nested.md").is_file()

        md_content = (tmp_path / "data.md").read_text(encoding="utf-8")
        assert "## 매출" in md_content and "## 현황" in md_content, md_content

        # pyhwpx가 없는 환경(예: 이 CI)이라면 HWP는 변환되지 않고 오류로 남는다.
        # pyhwpx가 있는 환경(Windows)이라면 이 assert는 건너뛴다.
        from folder_convert import Hwp as _Hwp
        if _Hwp is None:
            assert result["pdf_converted_count"] == 0, result
            assert any(e["type"] == "hwp_engine_missing" for e in result["errors"]), result

        print("[TEST 4/4] 폴더 일괄 변환 (XLSX -> Markdown 성공, HWP 엔진 부재 시 정상 보고): PASS")

    print("=" * 70)
    print("모든 테스트 통과 (ALL 4 TESTS PASSED)")
    print("=" * 70)


if __name__ == "__main__":
    run_all_tests()
